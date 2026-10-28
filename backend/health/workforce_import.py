"""Explicit, non-destructive, two-person roster import workflow.

Default environment gate WORKFORCE_IMPORT_ENABLED=false. Never stores workbook
bytes or raw preview employee identifiers. Existing employees are never updated.
"""
from datetime import timedelta
import hashlib
import io
import json
import os
import re
import zipfile

from django.core.exceptions import ValidationError as DjangoValidationError
from django.core.validators import validate_email
from django.db import IntegrityError, transaction
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework import permissions, status
from rest_framework.parsers import MultiPartParser, FormParser
from rest_framework.response import Response
from rest_framework.views import APIView

from .models import (
    AuditLog, Employee, HealthCenter, WorkforceCenterMapping,
    WorkforceImportPlan, WorkforceImportCreatedEmployee,
)
from .workforce_preview import (
    EXPECTED_COLUMNS, MAX_ROWS, MAX_UNCOMPRESSED, clean, norm, group_fingerprint
)

MAX_FILE_BYTES = 20 * 1024 * 1024
PLAN_LIFETIME = timedelta(hours=24)


def workbook_from_upload(upload):
    if not upload or not upload.name.lower().endswith('.xlsx'):
        raise ValueError('An original .xlsx workforce roster is required.')
    if upload.size > MAX_FILE_BYTES:
        raise ValueError('Workforce roster exceeds the 20 MB limit.')
    upload.seek(0)
    content = upload.read(MAX_FILE_BYTES + 1)
    if len(content) > MAX_FILE_BYTES or not content:
        raise ValueError('Invalid workbook size.')
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED:
                raise ValueError('Uncompressed workbook exceeds the limit.')
            if any(part.filename.startswith('/') or '..' in part.filename.split('/')
                   for part in archive.infolist()):
                raise ValueError('Invalid workbook content paths.')
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (zipfile.BadZipFile, ValueError) as exc:
        raise ValueError('Workbook is not a safe, valid XLSX file.') from exc
    except Exception as exc:
        raise ValueError('Unable to open the submitted XLSX file.') from exc
    return content, book


def mapping_snapshot():
    mappings = list(WorkforceCenterMapping.objects.filter(is_active=True)
        .select_related('health_center').order_by('fingerprint'))
    material = [
        [m.fingerprint, m.health_center_id, m.health_center.is_active,
         m.approved_at.isoformat() if m.approved_at else '']
        for m in mappings
    ]
    checksum = hashlib.sha256(json.dumps(
        material, ensure_ascii=False, separators=(',', ':')).encode('utf-8')).hexdigest()
    return {m.fingerprint: m for m in mappings}, checksum


def evaluate_import(upload):
    """Return short-lived in-memory candidate rows and non-identifying counts."""
    content, book = workbook_from_upload(upload)
    try:
        sheet = book.active
        if sheet.max_row and sheet.max_row > MAX_ROWS + 10:
            raise ValueError('Workbook row count exceeds the safe limit.')
        headers = [norm(val) for val in next(
            sheet.iter_rows(min_row=1, max_row=1, max_col=19, values_only=True))]
        if headers != [norm(col) for col in EXPECTED_COLUMNS]:
            raise ValueError('Only the original 19-column (A:S) workforce roster is supported.')

        mappings, mapping_hash = mapping_snapshot()
        existing = {
            item['national_id']: item for item in Employee.objects.values(
                'national_id', 'employee_number', 'email')
        }
        all_employee_numbers = {
            clean(item['employee_number']): key
            for key, item in existing.items() if item['employee_number']
        }
        all_emails = {
            norm(item['email']): key
            for key, item in existing.items() if item['email']
        }
        seen_id = set()
        seen_numbers = {}
        seen_emails = {}
        rows = []
        blocked_rows = 0
        missing_mapping_rows = 0
        existing_rows = 0
        new_rows = 0
        special_status_rows = 0
        errors = {}
        group_count = set()

        for number, row in enumerate(sheet.iter_rows(
            min_row=2, max_col=19, values_only=True), start=2
        ):
            values = [clean(value) for value in row]
            if not any(values):
                continue
            if (existing_rows + new_rows) >= MAX_ROWS:
                raise ValueError('Workbook has too many employee rows.')
            organization, org_type = values[0], values[1]
            identity = re.sub(r'\D', '', values[7])
            employee_number = values[6]
            name = values[8]
            gender_raw = norm(values[9])
            phone = re.sub(r'[^0-9+]', '', values[12])
            email = norm(values[13])
            source_status = values[14]
            source_appointment_date = values[15]
            job_title = values[17]
            group_key = group_fingerprint(organization, org_type)
            group_count.add(group_key)
            mapping = mappings.get(group_key)
            issues = []
            if not re.fullmatch(r'\d{10}', identity):
                issues.append('invalid_identity')
            if not name or len(name) > 200:
                issues.append('invalid_name')
            if not employee_number or len(employee_number) > 40:
                issues.append('invalid_employee_number')
            if len(job_title) > 150:
                issues.append('oversized_job_title')
            if len(source_status) > 120 or len(source_appointment_date) > 30:
                issues.append('oversized_source_fields')
            if phone and not (9 <= len(re.sub(r'\D', '', phone)) <= 15):
                issues.append('invalid_mobile')
            if email:
                try:
                    validate_email(email)
                except DjangoValidationError:
                    issues.append('invalid_email')
            if identity in seen_id:
                issues.append('duplicate_identity')
            if employee_number in seen_numbers and seen_numbers[employee_number] != identity:
                issues.append('duplicate_employee_number')
            if email and email in seen_emails and seen_emails[email] != identity:
                issues.append('duplicate_email')
            if identity not in existing:
                if employee_number in all_employee_numbers and all_employee_numbers[employee_number] != identity:
                    issues.append('existing_employee_number_conflict')
                if email and email in all_emails and all_emails[email] != identity:
                    issues.append('existing_email_conflict')
            if gender_raw in ('ذكر', 'male', 'm'):
                gender = 'male'
            elif gender_raw in ('انثى', 'أنثى', 'female', 'f'):
                gender = 'female'
            else:
                gender = ''
                issues.append('invalid_gender')
            if not organization or not org_type or not mapping or not mapping.health_center.is_active:
                issues.append('unapproved_center')
                missing_mapping_rows += 1
            if source_status != 'تعيين نشط':
                special_status_rows += 1

            if issues:
                blocked_rows += 1
                for issue in set(issues):
                    errors[issue] = errors.get(issue, 0) + 1
            if identity in existing:
                existing_rows += 1
            else:
                new_rows += 1
            seen_id.add(identity)
            if employee_number:
                seen_numbers[employee_number] = identity
            if email:
                seen_emails[email] = identity
            rows.append({
                'row': number, 'national_id': identity, 'name': name,
                'employee_number': employee_number, 'gender': gender,
                'email': email, 'mobile': phone,
                'health_center_id': mapping.health_center_id if mapping else None,
                'job_title': job_title,
                'source_status': source_status,
                'source_appointment_date': source_appointment_date,
                'existing': identity in existing,
                'issues': issues,
            })
        summary = {
            'total_rows': len(rows),
            'existing_unchanged': existing_rows,
            'new_candidates': new_rows,
            'blocked_rows': blocked_rows,
            'unapproved_center_rows': missing_mapping_rows,
            'other_source_status_rows': special_status_rows,
            'organization_groups': len(group_count),
            'issue_counts': errors,
            'ready': blocked_rows == 0 and bool(rows),
            'medical_records_modified': 0,
        }
        return {
            'file_sha256': hashlib.sha256(content).hexdigest(),
            'mapping_sha256': mapping_hash,
            'summary': summary,
            'rows': rows,
        }
    finally:
        book.close()


def summary_still_matches(plan, scan):
    return (
        plan.file_sha256 == scan['file_sha256'] and
        plan.mapping_sha256 == scan['mapping_sha256'] and
        plan.summary == scan['summary'] and
        scan['summary']['ready']
    )


def plan_expired(plan):
    return timezone.now() > plan.created_at + PLAN_LIFETIME


class WorkforceImportPlanView(APIView):
    permission_classes = [permissions.IsAdminUser]
    parser_classes = [FormParser, MultiPartParser]

    def get(self, request):
        recent = WorkforceImportPlan.objects.select_related(
            'created_by', 'approved_by').order_by('-created_at')[:30]
        return Response({'plans': [{
            'id': item.pk, 'state': item.state,
            'summary': item.summary,
            'file_sha256_prefix': item.file_sha256[:12],
            'created_by': item.created_by.username,
            'approved_by': item.approved_by.username if item.approved_by else None,
            'created_at': item.created_at.isoformat(),
            'expires_at': (item.created_at + PLAN_LIFETIME).isoformat(),
            'expired': plan_expired(item),
        } for item in recent]})

    def post(self, request):
        try:
            scan = evaluate_import(request.FILES.get('file'))
        except ValueError as exc:
            return Response({'detail': str(exc)}, status=status.HTTP_400_BAD_REQUEST)
        plan = WorkforceImportPlan.objects.create(
            file_sha256=scan['file_sha256'], mapping_sha256=scan['mapping_sha256'],
            summary=scan['summary'], created_by=request.user,
        )
        AuditLog.objects.create(
            user=str(request.user), action='workforce_import_plan_created',
            model_name='WorkforceImportPlan', record_id=str(plan.pk),
        )
        return Response({
            'plan_id': plan.pk, 'state': plan.state, 'summary': plan.summary,
            'file_sha256_prefix': plan.file_sha256[:12],
            'expires_at': (plan.created_at + PLAN_LIFETIME).isoformat(),
            'requires_second_approver': True, 'workbook_stored': False,
            'import_enabled': False,  # Never enable via a preview response.
        }, status=status.HTTP_201_CREATED)


class WorkforceImportApproveView(APIView):
    permission_classes = [permissions.IsAdminUser]

    def post(self, request, pk):
        approval_reference = clean(request.data.get('approval_reference'))
        backup_reference = clean(request.data.get('backup_reference'))
        restore_reference = clean(request.data.get('restore_verification_reference'))
        if any(not 4 <= len(item) <= 255 for item in
               (approval_reference, backup_reference, restore_reference)):
            return Response({'detail': 'Approval, backup, and restore references are required.'},
                            status=status.HTTP_400_BAD_REQUEST)
        with transaction.atomic():
            plan = WorkforceImportPlan.objects.select_for_update().filter(pk=pk).first()
            if not plan:
                return Response({'detail': 'Import plan was not found.'}, status=404)
            if plan.state != WorkforceImportPlan.State.DRAFT or plan_expired(plan):
                return Response({'detail': 'Plan is no longer valid or has expired.'}, status=400)
            if plan.created_by_id == request.user.pk:
                return Response({'detail': 'A different administrator must approve the plan.'}, status=403)
            if not plan.summary.get('ready'):
                return Response({'detail': 'Fix all invalid rows and mapping approvals first.'}, status=400)
            _, current_mapping_hash = mapping_snapshot()
            if plan.mapping_sha256 != current_mapping_hash:
                return Response({'detail': 'Center mappings changed; upload again for a new plan.'}, status=409)
            plan.state = WorkforceImportPlan.State.APPROVED
            plan.approved_by = request.user
            plan.approval_reference = approval_reference
            plan.backup_reference = backup_reference
            plan.restore_verification_reference = restore_reference
            plan.approved_at = timezone.now()
            plan.save(update_fields=[
                'state', 'approved_by', 'approval_reference', 'backup_reference',
                'restore_verification_reference', 'approved_at',
            ])
            AuditLog.objects.create(
                user=str(request.user), action='workforce_import_approved',
                model_name='WorkforceImportPlan', record_id=str(plan.pk),
            )
        return Response({
            'plan_id': plan.pk, 'state': plan.state,
            'note': 'References are administrator attestations, not independently verified backups.',
        })


class WorkforceImportCommitView(APIView):
    permission_classes = [permissions.IsAdminUser]
    parser_classes = [FormParser, MultiPartParser]

    def post(self, request, pk):
        if os.getenv('WORKFORCE_IMPORT_ENABLED', 'false').lower() != 'true':
            return Response(
                {'detail': 'Import is disabled until backup/restore testing and release authorization.'},
                status=status.HTTP_423_LOCKED,
            )
        if not request.user.is_superuser:
            return Response({'detail': 'Final import requires a superuser.'}, status=403)
        if clean(request.data.get('confirm')) != 'IMPORT_NEW_EMPLOYEES_ONLY':
            return Response({'detail': 'Explicit final confirmation phrase is required.'}, status=400)
        try:
            scan = evaluate_import(request.FILES.get('file'))
        except ValueError as exc:
            return Response({'detail': str(exc)}, status=400)

        with transaction.atomic():
            plan = WorkforceImportPlan.objects.select_for_update().filter(pk=pk).first()
            if not plan:
                return Response({'detail': 'Import plan was not found.'}, status=404)
            if plan.state != WorkforceImportPlan.State.APPROVED or plan_expired(plan):
                return Response({'detail': 'Approved plan expired or is not available.'}, status=409)
            if not plan.approved_by_id or plan.approved_by_id == plan.created_by_id:
                return Response({'detail': 'Missing independent approval.'}, status=409)
            if not all((plan.approval_reference, plan.backup_reference,
                        plan.restore_verification_reference)):
                return Response({'detail': 'Backup and restore references must be recorded.'}, status=409)
            if not summary_still_matches(plan, scan):
                return Response({'detail': 'File, database or center mappings changed. Prepare a new plan.'},
                                status=409)
            created = 0
            for row in scan['rows']:
                if row['existing']:
                    continue
                # Abort all writes on a concurrent identity/job/email collision.
                employee = Employee.objects.create(
                    national_id=row['national_id'],
                    name=row['name'], employee_number=row['employee_number'],
                    health_center_id=row['health_center_id'],
                    gender=row['gender'], email=row['email'] or None,
                    mobile=row['mobile'], job_title=row['job_title'],
                    roster_source_status=row['source_status'],
                    roster_hijri_appointment_date=row['source_appointment_date'],
                    periodic_exam_status='unknown',
                    vaccination_status='unknown', risk_level='unknown',
                )
                WorkforceImportCreatedEmployee.objects.create(
                    plan=plan, employee=employee, source_row=row['row'])
                created += 1
            if created != plan.summary['new_candidates']:
                # Include repeated/collision checks in the single rollback boundary.
                raise ValueError('Candidate count changed; transaction is rolled back.')
            plan.state = WorkforceImportPlan.State.COMMITTED
            plan.committed_at = timezone.now()
            plan.save(update_fields=['state', 'committed_at'])
            AuditLog.objects.create(
                user=str(request.user), action='workforce_import_committed',
                model_name='WorkforceImportPlan', record_id=str(plan.pk),
            )
        return Response({
            'plan_id': plan.pk, 'state': plan.state, 'new_employees_created': created,
            'existing_employees_unchanged': plan.summary['existing_unchanged'],
            'medical_records_modified': 0,
        })
