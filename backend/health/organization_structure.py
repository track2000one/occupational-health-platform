"""Organization matching-file intake and approval. No employee writes or implicit approvals."""
from collections import Counter
from hashlib import sha256
from io import BytesIO
import re
import zipfile

from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework import permissions, status
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.response import Response
from rest_framework.views import APIView

from .models import AuditLog, HealthCenter, OrganizationUnit, OrganizationUnitApprovalLog

MAX_FILE_SIZE = 5 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
MAX_ROWS_PER_SHEET = 250
IMPORT_MARKER = 'IMPORT_ORGANIZATION_DRAFTS_ONLY'
EXPECTED_SHEETS = ('المراكز المقترحة', 'وحدات المراكز', 'جهات مركزية')
STAGING_PROJECT = 'cf4b2162-5236-4f92-b010-165f505fe2c2'


def clean(value):
    return re.sub(r'\s+', ' ', str(value or '').strip())


def unique_sector_code(name):
    return 'TMP-S-' + sha256(clean(name).encode('utf-8')).hexdigest()[:12].upper()


def read_matching_workbook(upload):
    """Read only mapped reference sheets. Source workbook bytes are never retained."""
    if not upload or not upload.name.lower().endswith('.xlsx') or upload.size > MAX_FILE_SIZE:
        raise ValueError('Select the organizational matching XLSX file (maximum 5 MB).')
    content = upload.read(MAX_FILE_SIZE + 1)
    if len(content) > MAX_FILE_SIZE:
        raise ValueError('The matching file is too large.')
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            if sum(part.file_size for part in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                raise ValueError('Uncompressed workbook is too large.')
        book = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (zipfile.BadZipFile, OSError, ValueError) as exc:
        raise ValueError('Invalid or unsafe XLSX workbook.') from exc
    try:
        if any(name not in book.sheetnames for name in EXPECTED_SHEETS):
            raise ValueError('The three matching sheets are required. Use the approved review template.')
        rows = {}
        for sheet_name in EXPECTED_SHEETS:
            sheet = book[sheet_name]
            if sheet.max_row > MAX_ROWS_PER_SHEET:
                raise ValueError(f'Too many rows in {sheet_name}.')
            values = list(sheet.iter_rows(min_row=4, max_col=14, values_only=True))
            if not values:
                raise ValueError(f'Missing headers in {sheet_name}.')
            rows[sheet_name] = [list(row) for row in values[1:] if clean(row[0])]
        plans = {}
        center_codes = set()
        sector_codes = {}
        for row in rows['المراكز المقترحة']:
            code, sector, name = [clean(row[idx]) for idx in (0, 1, 2)]
            if not re.fullmatch(r'TMP-C\d{3}', code) or not sector or not name:
                raise ValueError('Invalid center reference or incomplete sector/center name.')
            sector_code = sector_codes.setdefault(sector, unique_sector_code(sector))
            plans[sector_code] = {
                'code': sector_code, 'name': sector, 'kind': 'sector',
                'parent_code': None, 'source_path': sector,
                'source_sheet': 'المراكز المقترحة', 'source_rows': '',
            }
            if code in center_codes:
                raise ValueError('Duplicate center reference in the template.')
            center_codes.add(code)
            plans[code] = {
                'code': code, 'name': name, 'kind': 'center',
                'parent_code': sector_code, 'source_path': clean(row[8]),
                'source_sheet': 'المراكز المقترحة', 'source_rows': clean(row[4]),
            }
        for row in rows['وحدات المراكز']:
            code, name, center_ref = clean(row[0]), clean(row[3]), clean(row[10])
            if not re.fullmatch(r'TMP-U\d{3}', code) or not name or center_ref not in center_codes:
                raise ValueError('A unit has no valid center or reference.')
            if code in plans:
                raise ValueError('Duplicate organizational reference in the template.')
            plans[code] = {
                'code': code, 'name': name, 'kind': 'unit',
                'parent_code': center_ref, 'source_path': clean(row[7]),
                'source_sheet': 'وحدات المراكز', 'source_rows': clean(row[6]),
            }
        for row in rows['جهات مركزية']:
            code, name = clean(row[0]), clean(row[2])
            if not re.fullmatch(r'TMP-O\d{3}', code) or not name:
                raise ValueError('Invalid central organization reference.')
            if code in plans:
                raise ValueError('Duplicate central reference in the template.')
            plans[code] = {
                'code': code, 'name': name, 'kind': 'department',
                # The proposed parent is NOT a verified organizational relationship.
                'parent_code': None, 'source_path': clean(row[3]),
                'source_sheet': 'جهات مركزية', 'source_rows': clean(row[6]),
            }
        return list(plans.values())
    finally:
        book.close()


def payload(item):
    return {
        'id': item.pk, 'code': item.code, 'name': item.name, 'kind': item.kind,
        'parent_id': item.parent_id, 'parent_code': item.parent.code if item.parent_id else '',
        'parent_name': item.parent.name if item.parent_id else '',
        'health_center_id': item.health_center_id,
        'health_center_name': item.health_center.name if item.health_center_id else '',
        'source_path': item.source_path, 'source_rows': item.source_rows,
        'is_approved': item.is_approved, 'is_active': item.is_active,
        'approval_reference': item.approval_reference,
    }


class OrganizationDirectoryView(APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        # Non-admin users never see pending organizational proposals.
        units = OrganizationUnit.objects.select_related('parent', 'health_center').filter(is_active=True)
        if not request.user.is_staff:
            units = units.filter(is_approved=True)
        if request.query_params.get('approved') == '1':
            units = units.filter(is_approved=True)
        items = list(units.order_by('kind', 'name', 'pk')[:1000])
        return Response({
            'results': [payload(u) for u in items],
            'summary': {
                'total': len(items),
                'approved': sum(u.is_approved for u in items),
                'pending': sum(not u.is_approved for u in items),
                'by_kind': dict(Counter(u.kind for u in items)),
            },
        })


class OrganizationImportView(APIView):
    permission_classes = [permissions.IsAdminUser]
    parser_classes = [MultiPartParser, FormParser]

    def post(self, request):
        try:
            plans = read_matching_workbook(request.FILES.get('file'))
        except (ValueError, TypeError) as exc:
            return Response({'detail': str(exc)}, status=status.HTTP_400_BAD_REQUEST)

        counts = Counter(p['kind'] for p in plans)
        results = {'total': len(plans), 'kinds': dict(counts),
                   'preview': plans[:12], 'created': 0,
                   'mode': 'preview', 'notice': 'No approvals or employee assignments were changed.'}
        if clean(request.data.get('confirm')) != IMPORT_MARKER:
            return Response(results)
        # Data owner has only approved preparation/review in the isolated staging project.
        import os
        if (os.getenv('STAGING_MODE', '').lower() != 'true'
                or os.getenv('RAILWAY_PROJECT_ID') != STAGING_PROJECT):
            return Response({'detail': 'Draft import is restricted to isolated Staging.'},
                            status=status.HTTP_403_FORBIDDEN)
        if not request.user.is_superuser:
            return Response({'detail': 'A staging superuser must import draft organizational data.'},
                            status=status.HTTP_403_FORBIDDEN)
        with transaction.atomic():
            for plan in plans:
                prev = OrganizationUnit.objects.filter(code=plan['code']).first()
                if prev:
                    if (prev.name, prev.kind, prev.parent.code if prev.parent_id else '') != (
                        plan['name'], plan['kind'], plan['parent_code'] or ''):
                        return Response({'detail': f'Existing reference {plan["code"]} conflicts; no records changed.'},
                                        status=status.HTTP_409_CONFLICT)
            for plan in plans:
                parent = (OrganizationUnit.objects.get(code=plan['parent_code'])
                          if plan['parent_code'] else None)
                unit, created = OrganizationUnit.objects.get_or_create(
                    code=plan['code'],
                    defaults={
                        'name': plan['name'], 'kind': plan['kind'], 'parent': parent,
                        'source_path': plan['source_path'], 'source_sheet': plan['source_sheet'],
                        'source_rows': plan['source_rows'], 'is_approved': False,
                    },
                )
                results['created'] += int(created)
            AuditLog.objects.create(user=str(request.user), action='organization_draft_import',
                                    model_name='OrganizationUnit', record_id='staging')
        results['mode'] = 'draft'
        return Response(results, status=status.HTTP_201_CREATED)


class OrganizationUnitReviewView(APIView):
    permission_classes = [permissions.IsAdminUser]

    def post(self, request, pk):
        if not request.user.is_superuser:
            return Response({'detail': 'Organizational approval requires a superuser.'},
                            status=status.HTTP_403_FORBIDDEN)
        action = clean(request.data.get('action'))
        ref = clean(request.data.get('approval_reference'))
        with transaction.atomic():
            unit = OrganizationUnit.objects.select_for_update().filter(pk=pk).first()
            if not unit:
                return Response({'detail': 'Organization not found.'}, status=404)
            if unit.is_approved:
                return Response({'detail': 'Already approved. Revisions require a separate review.'}, status=409)
            if action == 'map-center':
                if unit.kind != 'center':
                    return Response({'detail': 'Only a center entry can be mapped to a health center.'}, status=400)
                center = HealthCenter.objects.filter(pk=request.data.get('center_id'), is_active=True).first()
                if not center:
                    return Response({'detail': 'Choose an active center.'}, status=400)
                unit.health_center = center
                unit.save(update_fields=['health_center', 'updated_at'])
            elif action == 'approve':
                if len(ref) < 3 or len(ref) > 255:
                    return Response({'detail': 'Enter an official approval reference (3–255 characters).'}, status=400)
                if unit.parent_id and not unit.parent.is_approved:
                    return Response({'detail': 'Approve the parent sector or center first.'}, status=409)
                if unit.kind == 'center' and not unit.health_center_id:
                    return Response({'detail': 'Map the center to an existing registered health center first.'}, status=409)
                if unit.kind == 'unit':
                    if unit.parent.kind != 'center' or not unit.parent.health_center_id:
                        return Response({'detail': 'The parent center must have an approved health-center mapping.'}, status=409)
                    unit.health_center_id = unit.parent.health_center_id
                unit.is_approved = True
                unit.approved_by = request.user
                unit.approved_at = timezone.now()
                unit.approval_reference = ref
                unit.save()
            else:
                return Response({'detail': 'Action must be map-center or approve.'}, status=400)
            OrganizationUnitApprovalLog.objects.create(
                unit=unit, action=action, approval_reference=ref, user=request.user)
            AuditLog.objects.create(
                user=str(request.user), action='organization_' + action,
                model_name='OrganizationUnit', record_id=str(unit.pk),
            )
        unit.refresh_from_db()
        return Response(payload(OrganizationUnit.objects.select_related('parent', 'health_center').get(pk=unit.pk)))
